import sys, os, copy
MEDIA, ASSETS, SRC, OUT, TMP = sys.argv[1:6]
sys.argv = ['video2.py', MEDIA, ASSETS]
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import video2 as V2
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from lxml import etree

os.makedirs(TMP, exist_ok=True)
NAVY, RED, BG2, TX2, A4, A5, WHITE = '272B60', 'E30613', 'F1F2F9', '4B4F6B', 'DBDDEC', '8E93B5', 'FFFFFF'

# ---------- stage images (same visuals as the video) ----------
def stage(name, fn, lt, dur, inset=None):
    st = Image.new('RGB', (V2.SW, V2.SH), (255, 255, 255)); fn(st, lt, dur)
    if inset:
        im = V2.fit(V2.asset(*inset, trim=True, maxside=1600), 520, 300)
        x, y = V2.SW - im.width - 30, V2.SH - im.height - 30
        from PIL import ImageDraw
        ImageDraw.Draw(st).rectangle((x - 6, y - 6, x + im.width + 6, y + im.height + 6), fill=(255, 255, 255), outline=(219, 221, 236), width=3)
        st.paste(im, (x, y))
        V2.caption(st, 'Synoptique de détection', x, y - 50, V2.F(22, True), V2.NAVY)
    p = os.path.join(TMP, name + '.png'); st.save(p); return p

b5, b7, b4 = V2.N5.beats, V2.N7.beats, V2.N4.beats
ST = {
    'n5a': stage('n5a', b5[0][3], 0, 8),
    'n5b': stage('n5b', b5[1][3], 12, 8),
    'n5c': stage('n5c', V2.p_doe(['doe5_1.png', 'doe5_3.png', 'doe5_4.png', 'doe5_13.png', 'doe5_16.png', 'doe5_17.png', 'doe5_18.png',
                                  'doe5_19.png', 'doe5_21.png', 'doe5_22.png', 'doe5_23.png', 'doe5_25.png', 'doe5_26.png', 'doe5_28.png'],
                                 'DOE RIA généré automatiquement', 37, 'doe5_1.png'), 8.45, 12),
    'n7a': stage('n7a', V2.p_doc('doe7_2.png', (0.22, 0.15, 0.78, 0.75), cap="Plan d'implantation et câblage"), 0, 8,
                 inset=('doe7_5.png', (0.0, 0.0, 0.55, 0.42))),
    'n7b': stage('n7b', V2.p_doe(['doe7_%d.png' % i for i in range(1, 14)], 'DOE détection généré en 1 clic', 13, 'doe7_1.png'), 3.65, 7.2),
    'n7c': stage('n7c', b7[2][3], 9, 13),
    'n4a': stage('n4a', b4[0][3], 12.95, 13),
    'n4b': stage('n4b', V2.p_doc('tech_2.png', None, cap='Plan technicien – remplacement'), 0, 13),
}

def cover_crop(name, box, w_in, h_in, out):
    im = V2.asset(name, box, trim=True, maxside=2400)
    W, H = int(w_in * 300), int(h_in * 300)
    V2._fit(im, W, H, cover=True).save(os.path.join(TMP, out)); return os.path.join(TMP, out)

# ---------- pptx helpers ----------
prs = Presentation(SRC)
LAY = [l for l in prs.slide_layouts if l.name == 'CONTENU'][0]
SHADOW = ('<a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:outerShdw blurRad="101600" dist="25400" dir="5400000" '
          'algn="bl" rotWithShape="0"><a:srgbClr val="272B60"><a:alpha val="12000"/></a:srgbClr></a:outerShdw></a:effectLst>')

def rgb(h): return RGBColor.from_string(h)

def box(sl, x, y, w, h, fill, line=None, round_=True, shadow=False, name=None, adj=0.08):
    s = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if round_: s.adjustments[0] = adj
    s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    if line: s.line.color.rgb = rgb(line); s.line.width = Pt(1)
    else: s.line.fill.background()
    if shadow: s._element.spPr.append(etree.fromstring(SHADOW))
    if name: s.name = name
    s.text_frame.text = ''
    return s

def oval(sl, x, y, d, fill):
    s = sl.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    s.fill.solid(); s.fill.fore_color.rgb = rgb(fill); s.line.fill.background(); return s

def text(sl, x, y, w, h, runs, size=14, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=None, name=None):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    paras = runs if isinstance(runs[0], list) else [runs]
    for i, pr in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align
        for (t, col, b, sz) in pr:
            r = p.add_run(); r.text = t; r.font.color.rgb = rgb(col); r.font.bold = b; r.font.size = Pt(sz or size)
            if spacing: r.font._element.set('spc', str(spacing))
    if name: tb.name = name
    return tb

def pill(sl, x, y, label, fill, w=None, h=0.34, size=11, fg=WHITE):
    w = w or (0.012 * size * len(label) + 0.4)
    s = box(sl, x, y, w, h, fill, adj=0.5)
    tf = s.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; r = p.add_run(); r.text = label
    r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = rgb(fg); r.font._element.set('spc', '100')
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    return s

def icon_circle(sl, x, y, d, fill, ic):
    oval(sl, x, y, d, fill)
    k = d * 0.52; sl.shapes.add_picture(os.path.join(MEDIA, ic), Inches(x + (d - k) / 2), Inches(y + (d - k) / 2), Inches(k), Inches(k))

def title(sl, prefix, rest):
    ph = [p for p in sl.placeholders if p.placeholder_format.idx == 101][0]
    p = ph.text_frame.paragraphs[0]
    for t, col in ((prefix, 'accent1'), (rest, 'tx1')):
        if not t: continue
        r = p.add_run(); r.text = t
        rp = r.font._element; rp.set('b', '1'); rp.set('sz', '3000')
        sf = etree.SubElement(rp, '{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill')
        etree.SubElement(sf, '{http://schemas.openxmlformats.org/drawingml/2006/main}schemeClr').set('val', col)

def notes(sl, txt): sl.notes_slide.notes_text_frame.text = txt

def bandeau(sl, y, txt):
    s = box(sl, 0.6, y, 12.13, 0.46, NAVY, adj=0.15)
    text(sl, 0.85, y, 11.7, 0.46, [(txt, WHITE, True, 14)], anchor=MSO_ANCHOR.MIDDLE)

# ---------- slide: constat unique ----------
def constat():
    sl = prs.slides.add_slide(LAY); title(sl, '', 'Le constat : les difficultés du terrain')
    pill(sl, 12.05, 0.33, 'AVANT', RED, w=0.95)
    cols = V2.Constat.COLS
    for i, (lab, imgs, bul) in enumerate(cols):
        x = 0.6 + 4.14 * i; w = 3.85
        b = box(sl, x, 1.12, w, 0.4, NAVY, adj=0.15)
        text(sl, x, 1.12, w, 0.4, [(lab, WHITE, True, 15)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, spacing=200)
        iw = (w - 0.1) / 2
        for j, (nm, bx, cap) in enumerate(imgs):
            p = cover_crop(nm, bx, iw, 2.25, f'c{i}{j}.png')
            pic = sl.shapes.add_picture(p, Inches(x + j * (iw + 0.1)), Inches(1.65), Inches(iw), Inches(2.25))
            pic.line.color.rgb = rgb(A4); pic.line.width = Pt(0.75)
            cw = min(iw - 0.12, 0.075 * len(cap) + 0.2)
            c = box(sl, x + j * (iw + 0.1) + 0.06, 1.65 + 2.25 - 0.38, cw, 0.3, RED, adj=0.2)
            text(sl, x + j * (iw + 0.1) + 0.06, 1.65 + 2.25 - 0.38, cw, 0.3, [(cap, WHITE, True, 9)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        for j, (ic, tx) in enumerate(bul):
            y = 4.08 + 0.6 * j
            icon_circle(sl, x, y + 0.03, 0.42, RED, ic)
            text(sl, x + 0.55, y, w - 0.55, 0.48, [(tx, NAVY, True, 13)], anchor=MSO_ANCHOR.MIDDLE)
    bandeau(sl, 6.0, 'Résultat : des plans refaits à chaque étape, des dossiers inégaux, du temps perdu')
    notes(sl, "Un seul constat pour les trois métiers.\n"
              "N5 : l'étude est approximative (couverture vérifiée à la cordelette) et le plan est redessiné sur PowerPoint ; la note de calcul passe par Fathom (4 000 €/an).\n"
              "N7 : l'étude est longue, faite à la main ; le DOE est long à réaliser ; le plan technicien reste approximatif.\n"
              "N4 : des dossiers non homogènes (fluo, gommettes selon le commercial) et un plan technicien approximatif.\n"
              "Conclusion : on refait les plans à chaque étape et on perd du temps.")
    return sl

# ---------- slides: solutions ----------
def sol(prefix, label, name, beats, k, img, note):
    sl = prs.slides.add_slide(LAY); title(sl, prefix, "  –  Avec l'outil")
    pill(sl, 12.05, 0.33, 'APRÈS', NAVY, w=0.95)
    pic = sl.shapes.add_picture(img, Inches(0.45), Inches(1.0), Inches(8.3), Inches(8.3 * V2.SH / V2.SW))
    pic.line.color.rgb = rgb(A4); pic.line.width = Pt(1)
    pic._element.spPr.append(etree.fromstring(SHADOW))
    x, w = 8.98, 3.9
    box(sl, x, 1.0, w, 0.95, NAVY, shadow=True, adj=0.12)
    pill(sl, x + 0.2, 1.1, label, RED, w=0.11 * len(label) * 0.62 + 0.4, h=0.28, size=9)
    text(sl, x + 0.2, 1.43, w - 0.35, 0.45, [(name, WHITE, True, 18 if len(name) < 20 else 15)], anchor=MSO_ANCHOR.MIDDLE)
    y = 2.12
    for i, (bt, ti, sub, fn) in enumerate(beats):
        if i == k:
            h = 1.35; box(sl, x, y, w, h, NAVY, shadow=True, adj=0.1)
            oval(sl, x + 0.18, y + 0.2, 0.42, RED)
            text(sl, x + 0.18, y + 0.2, 0.42, 0.42, [(str(i + 1), WHITE, True, 16)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            text(sl, x + 0.75, y + 0.16, w - 0.9, 1.1, [[(ti, WHITE, True, 17)], [(sub, A4, False, 12)]])
        else:
            done = i < k; h = 0.72
            box(sl, x, y, w, h, WHITE if done else BG2, line=A4, adj=0.15)
            oval(sl, x + 0.18, y + h / 2 - 0.18, 0.36, NAVY if done else A4)
            if done: sl.shapes.add_picture(os.path.join(MEDIA, 'image-11-2.png'), Inches(x + 0.26), Inches(y + h / 2 - 0.1), Inches(0.2), Inches(0.2))
            else: text(sl, x + 0.18, y + h / 2 - 0.18, 0.36, 0.36, [(str(i + 1), WHITE, True, 13)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            text(sl, x + 0.68, y, w - 0.8, h, [(ti, NAVY if done else A5, True, 14)], anchor=MSO_ANCHOR.MIDDLE)
        y += h + 0.12
    notes(sl, note)
    return sl

N5n = ["Plan propre, étude nickel : l'implantation se fait sur le plan du client mis à l'échelle ; la couverture des RIA est calculée, murs et obstacles compris, avec des alertes en direct. Fini la cordelette.",
       "Note de calcul intégrée (HydraRIA) : pertes de charge, pression au RIA le plus défavorisé, temps d'arrivée d'eau, et la vue isométrique du réseau. Plus besoin de Fathom (4 000 €/an).",
       "Le point fort : le DOE se génère tout seul. Réception, plan d'implantation, plan isométrique et note de calcul, fiches techniques : 37 pages réunies par l'outil."]
N7n = ["Plan propre, étude nickel : paramètres R7 par local, implantation des détecteurs, câblage et synoptique dans le même outil.",
       "Le DOE en 1 clic : le dossier complet (plans par niveau, synoptique, bilan de puissance) est produit directement à partir de l'étude.",
       "Le point fort : le fichier Excel de chiffrage Eurofeu se remplit automatiquement avec les quantités de l'étude. Le chiffrage est prêt sans ressaisie."]
N4n = ["Le dossier d'étude N4 Eurofeu se complète automatiquement : pochette, caractéristiques de l'établissement, audit de sécurité R4, relevé sur site. Le même rendu pour tous les commerciaux.",
       "Le plan technicien est propre et numéroté : les appareils à remplacer, à réformer ou à ajouter sont repérés sur le plan avec la liste correspondante. Le technicien sait exactement où intervenir."]

new = [constat()]
for k, key in enumerate(['n5a', 'n5b', 'n5c']): new.append(sol('N5 · RIA', 'N5 · RIA', 'Implantation RIA Pro + HydraRIA', b5, k, ST[key], N5n[k]))
for k, key in enumerate(['n7a', 'n7b', 'n7c']): new.append(sol('N7 · Détection', 'N7 · DÉTECTION', 'DetectR7', b7, k, ST[key], N7n[k]))
for k, key in enumerate(['n4a', 'n4b']): new.append(sol('N4 · Extincteurs', 'N4 · EXTINCTEURS', 'N4 Conceptor', b4, k, ST[key], N4n[k]))

# ---------- speaker notes on kept slides ----------
orig = list(prs.slides)[:20]
NOTES = {1: "Présentation des trois outils métier N4, N5 et N7 : ce qu'ils changent concrètement, avant / après.",
         2: "Quatre temps : le constat, les solutions, l'apport pour Eurofeu et deux options, puis la démonstration.",
         3: "Une idée née du terrain il y a trois ans, développée sur mon temps personnel, avec l'aide d'un cousin développeur et d'un beau-frère pour les tests. Aujourd'hui trois outils opérationnels.",
         8: "Les trois métiers ont les mêmes freins : du temps perdu, un rendu qui varie d'une personne à l'autre, et des outils génériques (PowerPoint, Excel, Fathom). D'où l'idée d'outils pensés pour nos métiers.",
         10: "La même logique pour les trois outils : on charge le plan du client, on le met à l'échelle avec une cote connue, on implante avec les règles APSAD intégrées, et l'outil produit les livrables.",
         15: "Six gains pour Eurofeu : une longueur d'avance, des dossiers homogènes, du temps gagné, des dossiers plus professionnels, des économies (Fathom) et une conformité mieux maîtrisée.",
         16: "Pour donner un ordre de grandeur : de tels modules métier représentent environ 100 000 €.",
         17: "Option 1 : l'achat de licences auprès de mon beau-frère, qui les commercialise. Le tarif dépend du nombre de licences, de l'exclusivité, des modules et des services.",
         18: "Option 2 : une prime exceptionnelle, une augmentation et une évolution de poste, en échange des outils et de leur développement en interne pour Eurofeu.",
         19: "Place à la démonstration des trois outils en conditions réelles."}
for n, t in NOTES.items(): notes(orig[n - 1], t)

# ---------- reorder: replace slides 5-7 by the constat, 11-13 by the 8 solution slides ----------
sldIdLst = prs.slides._sldIdLst
ids = list(sldIdLst)
o = ids[:20]; nw = ids[20:]
order = o[0:4] + [nw[0]] + o[7:10] + nw[1:9] + o[13:20]
drop = [o[4], o[5], o[6], o[10], o[11], o[12]]
for e in ids: sldIdLst.remove(e)
for e in order: sldIdLst.append(e)
for e in drop: prs.part.drop_rel(e.rId)
prs.save(OUT)
print('slides:', len(order))
